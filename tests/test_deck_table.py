"""«добавь слайд с таблицей характеристик» made a slide titled so, with no table:
the deck had no table at all. Now a slide's "table" becomes a native pptx table."""
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("F5_TEST_RUN", "1")
import slides
from pptx import Presentation

deck = slides.normalize_deck({"title": "iPhone vs Galaxy", "theme": "tech", "slides": [
    {"heading": "Intro", "bullets": ["Two flagships compared on price and screen"]},
    {"heading": "Specs", "bullets": ["Side by side"],
     "table": {"columns": ["", "iPhone 16", "Galaxy S25"],
               "rows": [["Screen", "6.1\"", "6.2\""], ["RAM", "8 GB"], ["Chip", "A18", "SD 8 Elite", "extra"]]}}]})
t = deck["slides"][1]["table"]
assert t["rows"][1] == ["RAM", "8 GB", ""] and len(t["rows"][2]) == 3, t
assert slides._clean_table({"columns": ["only one"], "rows": [["x"]]}) == {}
out = os.path.join(tempfile.mkdtemp(), "t.pptx")
slides.build_pptx(deck, out)
cells = [c.text for sl in Presentation(out).slides for sh in sl.shapes if sh.has_table
         for row in sh.table.rows for c in row.cells]
assert "Galaxy S25" in cells and "SD 8 Elite" in cells, cells
assert '"table"' in slides._PLAN_PROMPT and "table" in slides._EDIT_PROMPT
print("ok a slide's table is drawn as a real table")
assert slides._clean_chart({"labels": ["Оперативная память (ГБ)", "Ядра нейропроцессора (ед)"], "values": [12, 16]}) == {}
assert slides._clean_chart({"labels": ["iPhone (ГБ)", "Galaxy (ГБ)"], "values": [8, 12]})
print("ok a chart mixing units on one axis is dropped")
assert "When the request names an audience" in slides._PLAN_PROMPT
print("ok the planner writes for a named audience")
# «сделай светлую тему»: there was no light theme, every cover was a dark panel with white text.
assert slides.pick_theme({"title": "X"}, "презентация про кофе, светлая тема")["name"] == "light"
out = os.path.join(tempfile.mkdtemp(), "l.pptx")
slides.build_pptx(dict(deck, theme="light"), out)
cov = Presentation(out).slides[0]
runs = [r for sh in cov.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
assert runs and all(str(r.font.color.rgb) != "FFFFFF" for r in runs), [str(r.font.color.rgb) for r in runs]
print("ok a light theme has a light cover with dark lettering")
# A figure in none of the notes is the planner's invention: dropped, rounding allowed.
_d = {"slides": [{"heading": "Tesla", "bullets": ["Основана в 2003 году", "Основана в 2008 году инвесторами",
                                                   "Выручка 96,8 млрд $ в 2023", "Около 97 млрд $", "3 модели"],
                  "stats": [{"value": "1,8 млн", "label": "авто"}, {"value": "5 млн", "label": "авто"}]}]}
slides._drop_ungrounded_figures(_d, "Tesla was founded in 2003. Revenue in 2023 was $96.77 billion; 1.81 million cars delivered.")
assert _d["slides"][0]["bullets"] == ["Основана в 2003 году", "Выручка 96,8 млрд $ в 2023", "Около 97 млрд $", "3 модели"], _d
assert [s["value"] for s in _d["slides"][0]["stats"]] == ["1,8 млн"], _d
print("ok figures from no note are dropped")
# The figure is in the notes, the sentence is not: MiniCheck's verdict drops it.
import fact_check, llm as _llm
_fs, _cl = fact_check.support, _llm.call_llm_simple
fact_check.support = lambda claims, notes: [0.9 if "2023" in c and notes == "Revenue 2023" else 0.02 for c in claims]
_llm.call_llm_simple = lambda ctx, s, u, **k: '["Revenue 2023", "Revenue was 96.8 billion in 2023", "Revenue was 96.8 billion in 2021"]'
try:
    _d = {"slides": [{"heading": "T", "bullets": ["Выручка 96,8 млрд в 2023", "Выручка 96,8 млрд в 2021", "Три модели"]}]}
    slides._drop_unsupported_claims(None, _d, "Выручка 2023")
    assert _d["slides"][0]["bullets"] == ["Выручка 96,8 млрд в 2023", "Три модели"], _d
finally:
    fact_check.support, _llm.call_llm_simple = _fs, _cl
print("ok a misattributed figure is dropped")
