"""A deck must carry FIGURES, CHARTS and LINKS -- not a list of headings.

The decks the assistant produced were laid out correctly and said nothing: bullets
like "Modern era" and "The future", no numbers anywhere, no chart, and no way to
check any of it. The planner prompt only ever set CEILINGS ("at most 12 words"),
which a model satisfies perfectly by writing two.

So the floor is enforced here in code, and this suite is what makes the
enforcement real: the parsers that turn a sloppy plan into something drawable,
the audit that decides whether a repair round is worth a turn, and the builder
that actually puts a native chart and clickable sources into the file.

Offline: no LLM, no search, no GPU.
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import slides  # noqa: E402

_fail = 0
_pass = 0


def check(name, cond, detail=""):
    global _fail, _pass
    if cond:
        _pass += 1
        print("PASS  " + name)
    else:
        _fail += 1
        print("FAIL  " + name + ((": " + str(detail)) if detail else ""))
        if os.environ.get("PYTEST_CURRENT_TEST"):
            raise AssertionError(str(name) + ((": " + str(detail)) if detail else ""))


def head(t):
    print("=" * 66)
    print(t)
    print("=" * 66)


# ── numbers as models actually write them ────────────────────────────────────
head("A FIGURE THE MODEL WROTE LOOSELY IS STILL A FIGURE")

for raw, want in (("3 200", 3200.0), ("12,100", 12100.0), ("12100", 12100.0),
                  ("2,9", 2.9), (4100, 4100.0), ("54 %", 54.0),
                  ("220 hp", 220.0), ("-15", -15.0)):
    check("%r parses to %s" % (raw, want), slides._to_number(raw) == want,
          slides._to_number(raw))
check("prose with no digit is not a number", slides._to_number("many") is None)
check("True is not a quantity", slides._to_number(True) is None)


# ── the stat band ────────────────────────────────────────────────────────────
head("THE STAT BAND IS FIGURES OR IT IS NOTHING")

check("a well-formed band survives",
      len(slides._clean_stats([{"value": "1962", "label": "first"},
                               {"value": "220 hp", "label": "power"}])) == 2)
check("a lone stat is dropped rather than drawn alone",
      slides._clean_stats([{"value": "1962", "label": "first"}]) == [])
check("a value with no digit is not a figure",
      slides._clean_stats([{"value": "many", "label": "a"},
                           {"value": "lots", "label": "b"}]) == [])
# Real figures, not "0".."8": a bare single digit is a label, not a
# measurement, and is rejected before the cap is ever reached.
check("the band is capped at four",
      len(slides._clean_stats([{"value": "%d т" % i, "label": "x"}
                               for i in range(9)])) == slides.MAX_STATS)


# ── charts ───────────────────────────────────────────────────────────────────
head("A CHART WE CANNOT DRAW MUST NOT REACH THE SLIDE")

good = slides._clean_chart({"kind": "bar", "labels": ["a", "b", "c"],
                            "values": ["1 200", 2400, "3,600"], "unit": "шт."})
check("units and separators are stripped to plain numbers",
      good.get("values") == [1200.0, 2400.0, 3600.0], good)
check("the unit is kept beside the numbers, not inside them",
      good.get("unit") == "шт.")

ragged = slides._clean_chart({"labels": ["a", "b", "c"], "values": [1, 2]})
check("ragged label/value lists are truncated to the pairs that exist",
      ragged.get("labels") == ["a", "b"] and ragged.get("values") == [1.0, 2.0],
      ragged)
check("a single point is not a chart",
      slides._clean_chart({"labels": ["a"], "values": [5]}) == {})
check("an all-zero series is not a chart",
      slides._clean_chart({"labels": ["a", "b"], "values": [0, 0]}) == {})
check("an unreadable value costs its point, not the chart",
      slides._clean_chart({"labels": ["a", "b", "c"],
                           "values": [1, "n/a", 3]}).get("values") == [1.0, 3.0])
check("a pie of negatives falls back to bars",
      slides._clean_chart({"kind": "pie", "labels": ["a", "b"],
                           "values": [-1, 5]}).get("kind") == "bar")
check("an unknown kind falls back to bars",
      slides._clean_chart({"kind": "radar", "labels": ["a", "b"],
                           "values": [1, 2]}).get("kind") == "bar")
check("a chart source must be a real URL",
      slides._clean_chart({"labels": ["a", "b"], "values": [1, 2],
                           "source": "see Wikipedia"}).get("source") == "")
check("a chart is capped at twelve points",
      len(slides._clean_chart({"labels": [str(i) for i in range(30)],
                               "values": list(range(1, 31))})["values"])
      == slides.MAX_CHART_POINTS)


# ── sources ──────────────────────────────────────────────────────────────────
head("A SOURCE IS A LINK SOMEBODY CAN OPEN")

got = slides._clean_sources([
    {"title": "One", "url": "https://example.org/a"},
    {"title": "Dup", "url": "HTTPS://EXAMPLE.ORG/A"},
    {"title": "Not a link", "url": "see Wikipedia"},
    {"title": "Bare domain", "url": "example.org"},
    "https://example.org/b",
])
check("only absolute http(s) links survive", [s["url"] for s in got]
      == ["https://example.org/a", "https://example.org/b"], got)
check("a bare URL gets itself as its title",
      got[1]["title"] == "https://example.org/b")


# ── the audit ────────────────────────────────────────────────────────────────
head("THE AUDIT MEASURES WHAT THE PLANNER DID, NOT WHAT IT WAS ASKED")

thin = slides.normalize_deck({"title": "T", "slides": [
    {"heading": "Modern era", "bullets": ["The future", "Progress"]},
    {"heading": "History", "bullets": ["Long ago"]}]})
a = slides.deck_facts_audit(thin)
check("a deck of headings fails the audit", not a["ok"], a)
check("and it fails for the right reason", a["charts"] == 0
      and a["slides_with_figures"] == 0, a)

rich = slides.normalize_deck({"title": "T", "sources": [
    {"title": "a", "url": "https://example.org/a"},
    {"title": "b", "url": "https://example.org/b"}], "slides": [
    {"heading": "One", "bullets": ["output rose to 12,100 units by 1985"],
     "chart": {"labels": ["1975", "1985"], "values": [9400, 12100]}},
    {"heading": "Two", "bullets": ["a claim"],
     "stats": [{"value": "54 %", "label": "share"},
               {"value": "2 300", "label": "per year"}]}]})
b = slides.deck_facts_audit(rich)
check("a deck with figures, a chart and links passes", b["ok"], b)
check("stats count as figures even when the bullets have none",
      b["slides_with_figures"] == 2, b)


# ── normalisation carries the new keys ───────────────────────────────────────
head("THE NEW KEYS SURVIVE NORMALISATION")

check("stats reach the slide", rich["slides"][1]["stats"][0]["value"] == "54 %")
check("charts reach the slide", rich["slides"][0]["chart"]["values"] == [9400.0, 12100.0])
lifted = slides.normalize_deck({"title": "T", "slides": [
    {"heading": "H", "bullets": ["b"],
     "sources": [{"title": "s", "url": "https://example.org/x"}]}]})
check("per-slide links are hoisted to the deck",
      [s["url"] for s in lifted["sources"]] == ["https://example.org/x"],
      lifted["sources"])
check("an empty plan still yields a slide with the new keys present",
      slides.normalize_deck({})["slides"][0]["chart"] == {})


# ── the file actually contains them ──────────────────────────────────────────
head("THE BUILT FILE CONTAINS A REAL CHART AND REAL HYPERLINKS")

deck = slides.normalize_deck({
    "title": "Кировец", "theme": "energy",
    "sources": [{"title": "Завод", "url": "https://example.org/k700"}],
    "slides": [
        {"heading": "Выпуск", "bullets": ["пик в 1985 году"],
         "stats": [{"value": "1962", "label": "старт"},
                   {"value": "220 л.с.", "label": "мощность"}],
         "chart": {"title": "Выпуск", "kind": "bar", "labels": ["1975", "1985"],
                   "values": ["9 400", "12,100"], "unit": "шт."}},
        {"heading": "Без цифр", "bullets": ["просто текст"]},
    ]})
out = os.path.join(tempfile.mkdtemp(), "deck.pptx")
written = slides.build_pptx(deck, out, topic="трактор")
check("the deck was written", bool(written) and os.path.exists(out), written)

if written:
    from pptx import Presentation
    prs = Presentation(out)
    n_charts = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
    check("the chart is a native pptx chart, not a picture of one", n_charts == 1,
          n_charts)
    if n_charts:
        ch = [sh.chart for s in prs.slides for sh in s.shapes if sh.has_chart][0]
        check("its numbers are the parsed ones",
              list(ch.plots[0].series[0].values) == [9400.0, 12100.0],
              list(ch.plots[0].series[0].values))
        check("its categories are the labels",
              list(ch.plots[0].categories) == ["1975", "1985"])
    # cover + 2 content + sources
    check("a sources slide was appended", len(prs.slides) == 4, len(prs.slides))
    xml = prs.slides[-1].shapes[-1]._element.xml
    check("the source line is a hyperlink, not grey text", "hlinkClick" in xml)
    def _first_text(slide):
        # The tinted panels are autoshapes and carry an EMPTY text frame, so
        # shapes[0] is not the heading -- take the first shape that says anything.
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                return sh.text_frame.text.strip()
        return ""

    heading = _first_text(prs.slides[-1])
    check("a Russian deck gets a Russian sources heading",
          heading == "Источники", heading)

    en = slides.normalize_deck({"title": "Tractors", "slides": [
        {"heading": "Output", "bullets": ["rose"]}],
        "sources": [{"title": "a", "url": "https://example.org/a"}]})
    out2 = os.path.join(os.path.dirname(out), "en.pptx")
    slides.build_pptx(en, out2)
    h2 = _first_text(Presentation(out2).slides[-1])
    check("an English deck gets an English one", h2 == "Sources", h2)

    # A bad chart must cost the chart, never the slide.
    broken = slides.normalize_deck({"title": "T", "slides": [
        {"heading": "H", "bullets": ["one", "two"],
         "chart": {"labels": ["a"], "values": ["n/a"]}}]})
    out3 = os.path.join(os.path.dirname(out), "broken.pptx")
    slides.build_pptx(broken, out3)
    p3 = Presentation(out3)
    texts = " ".join(sh.text_frame.text for s in p3.slides for sh in s.shapes
                     if sh.has_text_frame)
    check("an undrawable chart is dropped and the bullets survive",
          "one" in texts and "two" in texts, texts[:120])


# ── the picture must be of the SUBJECT, and findable ─────────────────────────
head("A SLIDE'S PICTURE SEARCH MUST NAME THE SUBJECT, BRIEFLY")

SUBJ = "История трактора Кировец"

# Measured against the live image search, same subject three ways:
#   "...tractor working in a vast field during sunset wide shot Кировец"
#        -> 3 hits, all vecteezy/freepik stock art
#   "Кировец ..." (the same words, subject moved to the front)
#        -> 3 hits, the same stock art
#   "Кировец К-700 трактор"
#        -> 6 hits, every one an actual Kirovets, the factory's site included
# So the lever is the amount of scene prose, not where the subject sits.
_long = ("massive heavy duty agricultural tractor working in a vast field "
         "during sunset wide shot")
_q = slides._image_query(_long, SUBJ)
check("the subject leads the query", _q.startswith("Кировец"), _q)
check("the query is short", len(_q.split()) <= slides._IMAGE_QUERY_WORDS, _q)
for word in ("sunset", "wide", "shot", "vast", "massive", "working"):
    check("photography vocabulary is dropped (%s)" % word, word not in _q, _q)
for word in ("tractor", "field"):
    check("the planner's real subject words survive (%s)" % word, word in _q, _q)

check("a query that already names the subject does not repeat it",
      slides._image_query("Кировец на поле", SUBJ).count("Кировец") == 1,
      slides._image_query("Кировец на поле", SUBJ))
check("a model number counts as the subject",
      slides._image_query("K-700 tractor field", SUBJ).startswith("Кировец"),
      slides._image_query("K-700 tractor field", SUBJ))
check("an empty query still names the subject",
      slides._image_query("", SUBJ) == "Кировец", slides._image_query("", SUBJ))
# "Витамин D" searched as "Витамин" is a different subject. The anchor filter
# kept only words of five letters or more (plus anything with a digit), so the
# single character that names the subject was the one thing thrown away, and a
# clinical deck came back with no cover picture at all.
check("a one-letter designator survives the anchor filter",
      "D" in slides._anchor_terms("Витамин D: доказательная база"),
      slides._anchor_terms("Витамин D: доказательная база"))
check("...and reaches the search query",
      "D" in slides._image_query("витамин D солнце и капсулы",
                                 "Витамин D: дозировки").split(),
      slides._image_query("витамин D солнце и капсулы", "Витамин D: дозировки"))
check("a designator leading nothing is not invented",
      slides._anchor_terms("D") == [], slides._anchor_terms("D"))
check("a lone preposition is still not an anchor",
      "в" not in slides._anchor_terms("Рынок электросамокатов в России"),
      slides._anchor_terms("Рынок электросамокатов в России"))
check("a lowercase short word is not mistaken for a designator",
      "и" not in slides._anchor_terms("Спрос и предложение на рынке"),
      slides._anchor_terms("Спрос и предложение на рынке"))
check("«кадр» is photography vocabulary, not a subject",
      "кадр" not in slides._image_query("витамин D, широкий кадр", "Витамин D"),
      slides._image_query("витамин D, широкий кадр", "Витамин D"))

check("a subject with nothing identifying in it changes nothing",
      slides._image_query("city street", "История") == "city street")
check("an English deck is handled too",
      slides._image_query("solar panels on a roof during golden hour wide shot",
                          "Renewable energy in Germany").startswith("Renewable"),
      slides._image_query("solar panels on a roof during golden hour wide shot",
                          "Renewable energy in Germany"))
check("...and drops its scene words as well",
      "golden" not in slides._image_query(
          "solar panels on a roof during golden hour wide shot",
          "Renewable energy in Germany"))


# ── the repair round ─────────────────────────────────────────────────────────
head("A THIN PLAN BUYS ONE REWRITE, AND ONLY IF IT COMES BACK RICHER")

import json as _json  # noqa: E402
import llm as _L  # noqa: E402

THIN = _json.dumps({"title": "T", "slides": [
    {"heading": "Modern era", "bullets": ["The future"]},
    {"heading": "History", "bullets": ["Long ago"]}]})
RICH = _json.dumps({"title": "T", "slides": [
    {"heading": "Modern era", "bullets": ["output rose to 12,100 units by 1985"],
     "chart": {"labels": ["1975", "1985"], "values": [9400, 12100]}},
    {"heading": "History", "bullets": ["built from 1962"],
     "stats": [{"value": "1962", "label": "start"},
               {"value": "220 hp", "label": "power"}]}]})

_real_gather, _real_llm = slides.gather_facts, _L.call_llm_simple
# Both boundaries replaced: this suite never opens a socket.
slides.gather_facts = lambda ctx, topic: (
    "- note: output rose to 12100 units by 1985; built from 1962, 220 hp; 9400 units in 1975",
    [{"title": "a", "url": "https://example.org/a"},
     {"title": "b", "url": "https://example.org/b"}])


def _scripted(*replies):
    box = {"n": 0}

    def _f(ctx, sys_p, user_p, **kw):
        box["n"] += 1
        return replies[min(box["n"], len(replies)) - 1]
    return box, _f


try:
    box, fn = _scripted(THIN, RICH)
    _L.call_llm_simple = fn
    deck2 = slides.plan_deck(None, "трактор Кировец")
    a2 = slides.deck_facts_audit(deck2)
    check("a deck of headings costs exactly one extra call", box["n"] == 2, box["n"])
    check("the rewrite is kept", a2["charts"] == 1 and a2["slides_with_figures"] == 2, a2)
    check("links absent from both plans are backfilled from the search",
          [s["url"] for s in deck2["sources"]]
          == ["https://example.org/a", "https://example.org/b"], deck2["sources"])
    check("and the finished deck passes its own audit", a2["ok"], a2)

    # A rewrite that comes back THINNER is a regression, not a repair.
    box, fn = _scripted(RICH.replace('"charts"', '"x"'), THIN)
    _L.call_llm_simple = fn
    deck3 = slides.plan_deck(None, "трактор Кировец")
    check("a thinner rewrite is rejected and the first plan stands",
          slides.deck_facts_audit(deck3)["slides_with_figures"] == 2,
          slides.deck_facts_audit(deck3))

    # A deck that was already good must not pay for a second call.
    box, fn = _scripted(_json.dumps({**_json.loads(RICH), "sources": [
        {"title": "a", "url": "https://example.org/a"},
        {"title": "b", "url": "https://example.org/b"}]}))
    _L.call_llm_simple = fn
    slides.plan_deck(None, "трактор Кировец")
    check("a deck that already has figures, a chart and links costs one call",
          box["n"] == 1, box["n"])
finally:
    slides.gather_facts, _L.call_llm_simple = _real_gather, _real_llm

# ── a deck nobody planned is not a deck ──────────────────────────────────────
head("A PLANNER THAT NEVER ANSWERED MUST NOT YIELD A PRESENTATION")

# Measured with LM Studio unloaded: make_presentation returned a real path and
# an EMPTY error string, so the caller announced a finished presentation. The
# file was a cover, one slide whose only text was the topic, and a sources
# list.
_real_llm2 = _L.call_llm_simple
_real_gather2 = slides.gather_facts
slides.gather_facts = lambda ctx, topic: ("", [])
_L.call_llm_simple = lambda *a, **k: ""
try:
    _dead = slides.plan_deck(None, "История трактора Кировец", 6)
    check("the failure is remembered on the deck", _dead.get("planner_failed") is True,
          _dead.get("planner_failed"))
    _res = slides.make_presentation(None, "История трактора Кировец",
                                    tempfile.mkdtemp(), n_slides=6)
    check("no path is handed back", _res.get("path") == "", _res.get("path"))
    check("and the error says why", "planner" in (_res.get("error") or ""),
          _res.get("error"))
    check("no images were fetched for a deck that does not exist",
          _res.get("images") == {}, _res.get("images"))
finally:
    _L.call_llm_simple = _real_llm2
    slides.gather_facts = _real_gather2

# A deck the planner DID write is never marked failed, even a short one.
_box3, _fn3 = _scripted(_json.dumps({"title": "T", "slides": [
    {"heading": "One", "bullets": ["built from 1962"]}]}))
slides.gather_facts = lambda ctx, topic: ("", [])
_L.call_llm_simple = _fn3
try:
    _one = slides.plan_deck(None, "тема", 2)
finally:
    _L.call_llm_simple = _real_llm2
    slides.gather_facts = _real_gather2
check("a one-slide deck the planner really wrote is NOT called a failure",
      _one.get("planner_failed") is False, _one.get("planner_failed"))


# ── a bare domain is not a citation ──────────────────────────────────────────
head("A SOURCE MUST POINT AT THE PAGE, NOT AT THE SITE")

_COLLECTED = [{"title": "a", "url": "https://igrader.ru/rbt/stalnaya-legenda"},
              {"title": "b", "url": "https://news.ru/auto/k700-history"},
              {"title": "c", "url": "https://x.ru/one"},
              {"title": "d", "url": "https://x.ru/two"}]
# Measured on a live deck: the planner cited news.ru, igrader.ru and
# kirovets-ptz.com as bare roots while the search that fed it had returned the
# specific articles on those very domains.
_fixed, _n = slides.repair_sources(
    [{"title": "Эволюция", "url": "https://igrader.ru"},
     {"title": "История", "url": "https://news.ru"},
     {"title": "Вики", "url": "https://ru.wikipedia.org/wiki/K-700"},
     {"title": "Двое", "url": "https://x.ru"}], _COLLECTED)
check("a bare root becomes the page found on it",
      _fixed[0]["url"] == "https://igrader.ru/rbt/stalnaya-legenda", _fixed[0])
check("...for every such source", _fixed[1]["url"].endswith("k700-history"), _fixed[1])
check("the planner's title is kept", _fixed[0]["title"] == "Эволюция", _fixed[0])
check("a domain the search never returned is DROPPED, not printed",
      all("wikipedia" not in f["url"] for f in _fixed), _fixed)
check("a domain with two collected pages is kept but NOT guessed at",
      any(f["url"] == "https://x.ru" for f in _fixed), _fixed)
check("and the repairs are counted", _n == 2, _n)

# The worst case is not a bare root -- it is a PLAUSIBLE deep path. Seen live:
# the planner offered "drive2.ru/z/story/13456789/", a round story id that
# looks precise and leads nowhere.
_inv, _ninv = slides.repair_sources(
    [{"title": "invented", "url": "https://news.ru/z/story/13456789/"}], _COLLECTED)
check("an invented deep path is replaced by the page we read",
      _inv[0]["url"] == "https://news.ru/auto/k700-history", _inv)

_dup, _ = slides.repair_sources(
    [{"title": "root", "url": "https://igrader.ru"},
     {"title": "same page", "url": "https://igrader.ru/rbt/stalnaya-legenda"}],
    _COLLECTED)
check("two entries repairing to one page are printed once", len(_dup) == 1, _dup)
check("nothing collected means nothing changes",
      slides.repair_sources([{"title": "t", "url": "https://a.ru"}], []) ==
      ([{"title": "t", "url": "https://a.ru"}], 0))
check("no sources is not a crash", slides.repair_sources([], _COLLECTED) == ([], 0))
check("www and a trailing slash still match the domain",
      slides.repair_sources([{"title": "t", "url": "https://www.news.ru/"}],
                            _COLLECTED)[1] == 1)


# ── a stat must be a measurement, not a label ────────────────────────────────
head("AN ORDINAL IS NOT A FIGURE")

# Measured on a live deck: the planner produced `1 — первый колесный гигант`,
# an ordinal crowbarred into the value slot to satisfy "contains a digit", and
# it was set on the slide in 26pt beside a real 1961.
for good in ("1961", "12 т", "54 %", "220 л.с.", "2 300", "от 2 500 000 ₽"):
    check("%r is a figure" % good, slides._is_figure(good))
for bad in ("1", "3", "7", "нет", "", "первый", "24/7", "100%"):   # filler (sleep deck, live)
    check("%r is not" % bad, not slides._is_figure(bad))

_mixed = slides._clean_stats([{"value": "1961", "label": "начало"},
                              {"value": "1", "label": "первый гигант"}])
check("a band left with one real figure is dropped rather than shown alone",
      _mixed == [], _mixed)
_ok = slides._clean_stats([{"value": "1961", "label": "начало"},
                           {"value": "220 л.с.", "label": "мощность"}])
check("two real figures still make a band", len(_ok) == 2, _ok)


# ── a wrapped bullet hangs under its text ────────────────────────────────────
head("A WRAPPED BULLET DOES NOT RESTART AT THE MARGIN")

_dw = slides.normalize_deck({"title": "T", "slides": [
    {"heading": "H", "bullets": ["История началась весной 1961 года во время "
                                 "визита Никиты Хрущева в США"]}]})
_outw = os.path.join(_tf.mkdtemp(), "wrap.pptx") if "_tf" in dir() else None
import tempfile as _tf2
_outw = os.path.join(_tf2.mkdtemp(), "wrap.pptx")
slides.build_pptx(_dw, _outw)
from pptx import Presentation as _P2
_body = [sh for s in _P2(_outw).slides for sh in s.shapes
         if sh.has_text_frame and "История" in sh.text_frame.text]
check("the bullet paragraph exists", len(_body) == 1, len(_body))
_pPr = _body[0].text_frame.paragraphs[0]._p.find(
    "{http://schemas.openxmlformats.org/drawingml/2006/main}pPr")
check("it carries a left margin", _pPr is not None and _pPr.get("marL"), _pPr)
check("...and a NEGATIVE first-line indent, which is what makes it hang",
      _pPr is not None and int(_pPr.get("indent", "0")) < 0, _pPr.get("indent") if _pPr is not None else None)
check("the two cancel, so the glyph still starts at the margin",
      _pPr is not None and int(_pPr.get("marL")) == -int(_pPr.get("indent")),
      (_pPr.get("marL"), _pPr.get("indent")) if _pPr is not None else None)


# ── the chart may not sit on the text ────────────────────────────────────────
head("A CHART TAKES WHAT IS LEFT, NOT WHAT IT WANTS")

# Measured on a clinical deck: a stat band, four bullets and a 3in chart left
# the bullets 0.35in of column, and the last two ran underneath the chart title.
_B4 = ["Диапазон 20-30 нг/мл теперь считается нормальным и здоровым",
       "Ранее этот уровень классифицировался как недостаточность",
       "Рекомендуемый узкий диапазон для коррекции дефицита: 30-60 нг/мл",
       "Целевые значения в нмоль/л составляют 75-150 нмоль/л"]
check("four bullets are known to need more than an inch",
      slides._bullets_height(_B4, 11.5, 14, 9) > 1.0,
      slides._bullets_height(_B4, 11.5, 14, 9))
check("...and two need less than four",
      slides._bullets_height(_B4[:2], 11.5, 14, 9)
      < slides._bullets_height(_B4, 11.5, 14, 9))
check("the fitter uses the same estimate it is measured by",
      slides._fit_bullets(_B4, 11.5, 0.4)[0] == 14,
      slides._fit_bullets(_B4, 11.5, 0.4))

_dense = slides.normalize_deck({"title": "T", "theme": "academic", "slides": [
    {"heading": "Целевые уровни", "bullets": _B4,
     "stats": [{"value": "20-30", "label": "нг/мл"},
               {"value": "30-60", "label": "нг/мл при дефиците"}],
     "chart": {"title": "Диапазоны", "kind": "bar",
               "labels": ["норма", "цель"], "values": [20, 60]}}]})
_od = os.path.join(_tf3.mkdtemp() if "_tf3" in dir() else __import__("tempfile").mkdtemp(),
                   "dense.pptx")
slides.build_pptx(_dense, _od, topic="витамин")
from pptx import Presentation as _P4  # noqa: E402
from pptx.util import Emu as _Emu  # noqa: E402
_sl = _P4(_od).slides[1]
_texts = [sh for sh in _sl.shapes if sh.has_text_frame and "Диапазон 20-30" in sh.text_frame.text]
_charts = [sh for sh in _sl.shapes if sh.has_chart]
check("the bullets are on the slide", len(_texts) == 1, len(_texts))
if _texts and _charts:
    _b_top = _texts[0].top
    _need = slides._bullets_height(_B4, 11.5, 14, 9) * 914400
    _c_top = _charts[0].top
    check("the chart starts BELOW the room the bullets need",
          _c_top >= _b_top + _need, (_b_top, _need, _c_top))
elif _texts:
    check("a chart that cannot fit is dropped rather than overlapped", True)

# The reservation must not become a reason to drop the chart over a fraction of
# an inch: measured, this slide leaves it 1.74in and a short chart still reads.
check("a chart is kept when the leftover is short but usable",
      len(_charts) == 1, len(_charts))
if _charts:
    check("...and it is at least the minimum height",
          _charts[0].height / 914400.0 >= slides._MIN_CHART_IN,
          _charts[0].height / 914400.0)


# ── a chart wears the deck's colour, whatever its kind ───────────────────────
head("A LINE CARRIES ITS COLOUR ON THE LINE, NOT THE FILL")

# Seen on a slide: the bar chart came out in the deck's accent and the line
# chart in PowerPoint default blue, because both were coloured through the FILL
# and a line has none.
from pptx import Presentation as _P3  # noqa: E402
import tempfile as _tf3  # noqa: E402


def _one_chart(kind):
    _d = slides.normalize_deck({"title": "T", "theme": "energy", "slides": [
        {"heading": "H", "bullets": ["x"],
         "chart": {"title": "C", "kind": kind, "labels": ["1964", "1975"],
                   "values": [100, 100000]}}]})
    _o = os.path.join(_tf3.mkdtemp(), "c.pptx")
    slides.build_pptx(_d, _o, topic="трактор")
    _ch = [sh.chart for s in _P3(_o).slides for sh in s.shapes if sh.has_chart][0]
    return _ch.plots[0].series[0], "%06X" % slides.pick_theme(_d, "трактор")["accent"]


def _rgb_of(colour):
    """The colour as a hex string, or None.

    python-pptx raises AttributeError for .rgb on an unset colour, and a test
    that DIES on a regression instead of reporting it hides how many other
    checks would also have failed -- and reads as a crash rather than a verdict.
    """
    try:
        return str(colour.rgb)
    except Exception:
        return None


_ser, _acc = _one_chart("line")
check("a line series is coloured on its LINE",
      _rgb_of(_ser.format.line.color) == _acc, _rgb_of(_ser.format.line.color))
check("...and is thick enough to read from a seat at the back",
      _ser.format.line.width is not None)
_ser_b, _acc_b = _one_chart("bar")
check("a bar series is still coloured on its FILL",
      _rgb_of(_ser_b.format.fill.fore_color) == _acc_b,
      _rgb_of(_ser_b.format.fill.fore_color))


# ── links wear the deck's colour, not Office blue ────────────────────────────
head("A LINK IS PAINTED BY THE THEME, NOT BY THE RUN")

import tempfile as _tf  # noqa: E402
from pptx import Presentation as _P  # noqa: E402

_THEME_REL = ("http://schemas.openxmlformats.org/officeDocument/2006/"
              "relationships/theme")


def _theme_xml(path):
    prs = _P(path)
    part = prs.slide_masters[0].part.part_related_by(_THEME_REL)
    return part.blob.decode("utf-8")


_d = slides.normalize_deck({
    "title": "Кировец", "theme": "energy",
    "sources": [{"title": "Завод", "url": "https://example.org/k700"}],
    "slides": [{"heading": "Выпуск", "bullets": ["пик в 1985 году"]}]})
_out = os.path.join(_tf.mkdtemp(), "themed.pptx")
slides.build_pptx(_d, _out, topic="трактор")
_xml = _theme_xml(_out)
_accent = "%06X" % slides.pick_theme(_d, "трактор")["accent"]
check("the theme's link colour is the deck's accent",
      '<a:hlink><a:srgbClr val="%s"/>' % _accent in _xml, _xml[_xml.find("<a:hlink>"):][:60])
check("...and so is the FOLLOWED link colour",
      '<a:folHlink><a:srgbClr val="%s"/>' % _accent in _xml,
      _xml[_xml.find("<a:folHlink>"):][:60])
check("Office blue is gone", '<a:hlink><a:srgbClr val="0000FF"/>' not in _xml)
check("and so is the purple of a visited link", "800080" not in _xml)

# A different palette must produce a different link colour, or the repaint is
# writing a constant rather than the theme.
_d2 = slides.normalize_deck({
    "title": "Clinical trial design", "theme": "academic",
    "sources": [{"title": "A", "url": "https://example.org/a"}],
    "slides": [{"heading": "Method", "bullets": ["n = 240"]}]})
_out2 = os.path.join(os.path.dirname(_out), "themed2.pptx")
slides.build_pptx(_d2, _out2)
_accent2 = "%06X" % slides.pick_theme(_d2, "")["accent"]
check("a different theme yields a different link colour", _accent2 != _accent,
      (_accent, _accent2))
check("...and the file carries it",
      '<a:hlink><a:srgbClr val="%s"/>' % _accent2 in _theme_xml(_out2), _accent2)

print()
print("=" * 70)
print("THE BUILT FILE, NOT THE PLAN")
print("=" * 70)

# deck_facts_audit reads the PLAN. It reported two charts on a clinical deck
# whose file had none -- the chart reservation dropped both over a fraction of
# an inch, and the number meant to catch a thin deck was reading the wrong side
# of the build. pptx_audit opens the file instead.
#
# This is the exact slide that failed: a three-value stat band, four bullets,
# and a chart, in a body column 4.5in tall. Reproduced without the model.
_crowded = slides.normalize_deck({
    "title": "Витамин D: доказательная база", "theme": "academic",
    "sources": [{"title": "A", "url": "https://example.org/a"},
                {"title": "B", "url": "https://example.org/b"}],
    "slides": [{
        "heading": "Целевые уровни в крови",
        "stats": [{"value": "30 нг/мл", "label": "нижняя граница нормы"},
                  {"value": "60 нг/мл", "label": "целевой уровень"},
                  {"value": "100 нг/мл", "label": "верхний предел"}],
        "bullets": [
            "Дефицитом считают уровень 25(OH)D ниже 20 нг/мл",
            "Недостаточность — диапазон от 20 до 30 нг/мл",
            "Целевой коридор для большинства взрослых 30–60 нг/мл",
            "Свыше 100 нг/мл возрастает риск гиперкальциемии"],
        "chart": {"title": "Доля дефицита по возрасту", "kind": "bar",
                  "labels": ["18–30", "31–50", "51+"], "values": [42, 55, 71],
                  "unit": "%"}}]})
_outc = os.path.join(os.path.dirname(_out), "crowded.pptx")
slides.build_pptx(_crowded, _outc)
_audit = slides.pptx_audit(_outc)

check("the chart the plan promised is IN THE FILE", _audit["charts"] == 1, _audit)
check("...and nothing on the deck overlaps anything else",
      _audit["overlaps"] == [], _audit["overlaps"])
check("the file audit counts the slides it actually wrote",
      _audit["slides"] == 3, _audit)     # cover + content + sources

# A picture slide is the other half: the text column used to end 0.08in inside
# the photo panel. Every heading in three decks happened to be short enough to
# hide it.
from PIL import Image as _Im
_imgp = os.path.join(os.path.dirname(_out), "panel.png")
_Im.new("RGB", (800, 1200), (40, 90, 140)).save(_imgp)
_pic = slides.normalize_deck({
    "title": "История трактора", "theme": "energy",
    "slides": [{"heading": "Наследие Кировского завода и его конструкторов",
                "bullets": ["Производство началось в 1924 году",
                            "К 1937 году выпущено более 40 000 машин"],
                "image": "tractor"}]})
_outp = os.path.join(os.path.dirname(_out), "withpic.pptx")
slides.build_pptx(_pic, _outp, images={0: _imgp})
_pa = slides.pptx_audit(_outp)
check("the picture is placed", _pa["pictures"] == 1, _pa)
check("the text column does not run under the photo panel",
      _pa["overlaps"] == [], _pa["overlaps"])

# The detector must be able to SEE an overlap, or its silence means nothing.
_bad = slides.normalize_deck({
    "title": "T", "slides": [{"heading": "H", "bullets": ["x"]}]})
_outb = os.path.join(os.path.dirname(_out), "overlap_probe.pptx")
slides.build_pptx(_bad, _outb)
from pptx import Presentation as _P
from pptx.util import Inches as _I
_prs = _P(_outb)
_s = list(_prs.slides)[1]
_tb = _s.shapes.add_textbox(_I(1.0), _I(3.0), _I(3.0), _I(1.0))
_tb.text_frame.text = "deliberate collision"
_tb2 = _s.shapes.add_textbox(_I(2.0), _I(3.2), _I(3.0), _I(1.0))
_tb2.text_frame.text = "second box"
_prs.save(_outb)
check("a deliberate collision IS reported",
      slides.pptx_audit(_outb)["overlaps"], "detector is blind")

# Nothing may hang off the canvas: whatever is past the edge is simply lost.
check("no shape runs off the slide on a crowded deck",
      _audit["offslide"] == [], _audit["offslide"])
check("...nor on a picture slide", _pa["offslide"] == [], _pa["offslide"])
check("nothing overflows its box on a crowded deck",
      _audit["overflow"] == [], _audit["overflow"])

# A slide carrying nothing but its own heading is a defect the plan-side numbers
# cannot see. Found in a real deck in runtime/, from a run where the planner had
# failed: heading, footer, page number, and nothing else. The refusal upstream
# covers that cause; the audit names the symptom whatever causes it next time.
check("a real deck has no content-free slides", _audit["empty"] == [], _audit["empty"])
check("...nor does the picture one", _pa["empty"] == [], _pa["empty"])

_thin = slides.normalize_deck({"title": "T", "slides": [{"heading": "Только заголовок"}]})
_outt = os.path.join(os.path.dirname(_out), "thin.pptx")
slides.build_pptx(_thin, _outt)
check("a heading-only slide IS reported",
      slides.pptx_audit(_outt)["empty"] == [2],
      slides.pptx_audit(_outt)["empty"])
# The cover is a title slide BY DESIGN -- heading, subtitle, nothing else -- so
# it must never be reported. (The first version of this check asserted
# `2 not in [1]`, which is a tautology: it could not have failed.)
check("...and the cover is never counted as empty",
      1 not in slides.pptx_audit(_outt)["empty"],
      slides.pptx_audit(_outt)["empty"])

# A slide whose content is a CHART and nothing else is not empty.
_chart_only = slides.normalize_deck({
    "title": "T", "slides": [{"heading": "Только график",
                              "chart": {"title": "Выпуск", "kind": "bar",
                                        "labels": ["a", "b"], "values": [1, 2]}}]})
_outco = os.path.join(os.path.dirname(_out), "chartonly.pptx")
slides.build_pptx(_chart_only, _outco)
check("a chart counts as content", slides.pptx_audit(_outco)["empty"] == [],
      slides.pptx_audit(_outco)["empty"])

# The chart's unit was collected, validated -- and shown nowhere. It becomes the
# SERIES name, and bar and line charts have their legend switched off, so a
# chart of percentages rendered as bars of 42, 55 and 71 with no "%" anywhere on
# the slide. Seen on an exported clinical deck.
check("the unit reaches the chart title",
      slides._chart_title_with_unit({"title": "Доля дефицита", "unit": "%"})
      == "Доля дефицита, %")
check("a title that already says the unit is not doubled",
      slides._chart_title_with_unit({"title": "Доля рынка, %", "unit": "%"})
      == "Доля рынка, %")
check("no unit, no change",
      slides._chart_title_with_unit({"title": "Выпуск по годам", "unit": ""})
      == "Выпуск по годам")
check("a bare unit is not promoted to a title",
      slides._chart_title_with_unit({"title": "", "unit": "%"}) == "")
check("a trailing comma is not doubled",
      slides._chart_title_with_unit({"title": "Выпуск,", "unit": "шт"})
      == "Выпуск, шт")

# And it has to reach the FILE, not just the helper.
_ud = slides.normalize_deck({
    "title": "T", "slides": [{"heading": "H", "bullets": ["42% в 2024 году"],
                              "chart": {"title": "Доля дефицита", "kind": "bar",
                                        "labels": ["18-30", "31-50"],
                                        "values": [42, 55], "unit": "%"}}]})
_outu = os.path.join(os.path.dirname(_out), "unit.pptx")
slides.build_pptx(_ud, _outu)
from pptx import Presentation as _P2
_titles = [sh.chart.chart_title.text_frame.text
           for sl in _P2(_outu).slides for sh in sl.shapes
           if getattr(sh, "has_chart", False)]
check("the built chart carries its unit", _titles == ["Доля дефицита, %"], _titles)

# The headline figure was fixed at 26pt however long it was and however narrow
# the card. Caught by the FILE audit before anything was exported: "1 200 000 ₽"
# needed 1.19in of a 0.9in card on a slide that also carried a photograph, so
# three stats shared 6.8in instead of 11.5.
check("a long figure in a narrow card is set smaller",
      slides._fit_stat("1 200 000 ₽", 1.87) < slides._fit_stat("14,8%", 1.87),
      (slides._fit_stat("1 200 000 ₽", 1.87), slides._fit_stat("14,8%", 1.87)))
check("a short figure keeps the full size", slides._fit_stat("1962", 3.5) == 26)
check("...even in a narrow card", slides._fit_stat("14,8%", 1.87) == 26)
check("a figure never shrinks below its own label",
      slides._fit_stat("x" * 200, 1.0) > 11, slides._fit_stat("x" * 200, 1.0))
check("an empty value does not raise", slides._fit_stat("", 2.0) == 26)

_sd = slides.normalize_deck({
    "title": "T", "theme": "corporate",
    "slides": [{"heading": "Три цифры рядом с фотографией",
                "stats": [{"value": "1 200 000 ₽", "label": "стоимость владения"},
                          {"value": "14,8%", "label": "доля рынка"},
                          {"value": "3,2 года", "label": "срок окупаемости"}],
                "bullets": ["Доля рынка выросла с 9,4% за два года"],
                "image": "x"}]})
_outs = os.path.join(os.path.dirname(_out), "stats.pptx")
slides.build_pptx(_sd, _outs, images={0: _imgp})
_sa = slides.pptx_audit(_outs)
check("nothing in the stat band overflows its card", _sa["overflow"] == [],
      _sa["overflow"])
check("...and the band does not collide with anything", _sa["overlaps"] == [],
      _sa["overlaps"])

# The axes and gridlines were never styled at all, so PowerPoint drew them in
# its default near-black while every word on the deck is a soft grey. Third of
# the same family, after the line colour and the markers: the theme reached the
# series and stopped there.
_ad = slides.normalize_deck({
    "title": "T", "theme": "energy",
    "slides": [{"heading": "H", "bullets": ["разброс 2,4 раза"],
                "chart": {"title": "Парк", "kind": "bar",
                          "labels": ["a", "b", "c"], "values": [3, 2, 1],
                          "unit": "ед."}}]})
_outa = os.path.join(os.path.dirname(_out), "axes.pptx")
slides.build_pptx(_ad, _outa)
from pptx import Presentation as _P5
_ch = None
for _sl in _P5(_outa).slides:
    for _sh in _sl.shapes:
        if getattr(_sh, "has_chart", False):
            _ch = _sh.chart
_body_hex = "%06X" % slides.pick_theme(_ad, "")["body"]
check("the axis line takes the deck's body colour",
      str(_ch.category_axis.format.line.color.rgb) == _body_hex,
      str(_ch.category_axis.format.line.color.rgb))
check("...and so do the tick labels",
      str(_ch.value_axis.tick_labels.font.color.rgb) == _body_hex)
check("the gridlines are a TINT of it, not the same weight",
      str(_ch.value_axis.major_gridlines.format.line.color.rgb) != _body_hex,
      str(_ch.value_axis.major_gridlines.format.line.color.rgb))
check("_blend leaves a colour alone at t=0",
      slides._blend(0x3D342C, 0xFFFFFF, 0.0) == 0x3D342C)
check("_blend reaches the target at t=1",
      slides._blend(0x3D342C, 0xFFFFFF, 1.0) == 0xFFFFFF)
check("_blend clamps a nonsense t", slides._blend(0x000000, 0xFFFFFF, 5.0) == 0xFFFFFF)

# The line got the theme; its MARKERS did not. A warm orange line with default
# blue diamonds sitting on it, seen on the same exported deck -- one level down
# from the line-colour bug, and the same shape of mistake.
_ld = slides.normalize_deck({
    "title": "T", "theme": "energy",
    "slides": [{"heading": "H", "bullets": ["выпуск вырос в 3,2 раза"],
                "chart": {"title": "Выпуск", "kind": "line",
                          "labels": ["1964", "1970", "1975"],
                          "values": [1200, 3820, 9400], "unit": "шт."}}]})
_outl = os.path.join(os.path.dirname(_out), "line.pptx")
slides.build_pptx(_ld, _outl)
from pptx import Presentation as _P4
_ser = None
for _sl in _P4(_outl).slides:
    for _sh in _sl.shapes:
        if getattr(_sh, "has_chart", False):
            _ser = _sh.chart.plots[0].series[0]
_acc = "%06X" % slides.pick_theme(_ld, "")["accent"]
check("the line takes the theme accent",
      _ser is not None and str(_ser.format.line.color.rgb) == _acc,
      None if _ser is None else str(_ser.format.line.color.rgb))
check("...and so do its markers",
      str(_ser.marker.format.fill.fore_color.rgb) == _acc,
      str(_ser.marker.format.fill.fore_color.rgb))
check("the markers are not left at PowerPoint's blue",
      str(_ser.marker.format.fill.fore_color.rgb) != "4472C4")

# A pie was the one chart kind nobody had coloured. Bars and lines take the
# theme accent; the pie took PowerPoint's default blue/red/green, rendered on a
# warm brown deck. Same family as the line-colour bug: the theme reached some
# chart kinds and not others.
_pd = slides.normalize_deck({
    "title": "T", "theme": "energy",
    "slides": [{"heading": "H", "bullets": ["На три марки приходится 78% рынка"],
                "chart": {"title": "Доли рынка", "kind": "pie",
                          "labels": ["Кировец", "МТЗ", "Ростсельмаш", "Прочие"],
                          "values": [41, 22, 15, 22], "unit": "%"}}]})
_outp2 = os.path.join(os.path.dirname(_out), "pie.pptx")
slides.build_pptx(_pd, _outp2)
from pptx import Presentation as _P3
_pie = None
for _sl in _P3(_outp2).slides:
    for _sh in _sl.shapes:
        if getattr(_sh, "has_chart", False):
            _pie = _sh.chart
_accent_hex = "%06X" % slides.pick_theme(_pd, "")["accent"]
_slice_rgb = [str(p.format.fill.fore_color.rgb)
              for p in _pie.plots[0].series[0].points] if _pie else []

check("the pie is drawn at all", _pie is not None)
check("every slice is coloured explicitly", len(_slice_rgb) == 4, _slice_rgb)
check("the first slice IS the deck's accent", _slice_rgb[:1] == [_accent_hex],
      (_slice_rgb[:1], _accent_hex))
check("the slices are all different", len(set(_slice_rgb)) == 4, _slice_rgb)
check("none of them is PowerPoint's default blue",
      "4472C4" not in _slice_rgb, _slice_rgb)

# A pie has no axis, so unlabelled slices cannot be read at all.
check("the slices carry their values", _pie.plots[0].has_data_labels)

# The first attempt reached for chart.plots[0].points, which PiePlot does not
# have; the AttributeError went into an `except` and a still-default pie shipped
# while every check passed. So this asserts on the FILE, not on the helper.
_shades = slides._pie_shades(slides.pick_theme(_pd, "")["accent"], 4)
check("the shades run from the accent toward light",
      slides._is_light(_shades[-1]) and not slides._is_light(_shades[0]),
      [hex(c) for c in _shades])
check("a one-slice pie is just the accent",
      slides._pie_shades(0xD2691E, 1) == [0xD2691E])
check("a zero-slice pie does not raise", len(slides._pie_shades(0xD2691E, 0)) == 1)

# Sparse slides used to be as small as dense ones: the fitter could shrink but
# never grow, so two bullets filled the top fifth of a 4.5in column and left the
# rest of the slide blank. Raising the ceiling flat to 30pt then made SIX
# bullets 30pt too -- they fit, and a wall of huge type is a worse slide than an
# empty one. So the ceiling depends on how many bullets there are.
_two = ["Первая мысль с числом: 600 МЕ в сутки", "Вторая мысль: 125 мкг предел"]
_six = ["Дефицитом считают уровень 25(OH)D ниже 20 нг/мл"] * 6
check("two bullets are allowed to grow",
      slides._fit_bullets(_two, 11.5, 4.5)[0] >= 30,
      slides._fit_bullets(_two, 11.5, 4.5))
check("six bullets are NOT, even though they would fit",
      slides._fit_bullets(_six, 11.5, 4.5)[0] <= 22,
      slides._fit_bullets(_six, 11.5, 4.5))
check("the size falls as the count rises, without exception",
      [slides._fit_bullets(["Ровно одна короткая мысль"] * n, 11.5, 4.5)[0]
       for n in (2, 3, 4, 5, 6)] == sorted(
          [slides._fit_bullets(["Ровно одна короткая мысль"] * n, 11.5, 4.5)[0]
           for n in (2, 3, 4, 5, 6)], reverse=True),
      [slides._fit_bullets(["Ровно одна короткая мысль"] * n, 11.5, 4.5)[0]
       for n in (2, 3, 4, 5, 6)])
check("a narrow column still shrinks below its ceiling when it must",
      slides._fit_bullets(["очень длинная мысль " * 12], 3.0, 1.0)[0] <= 18,
      slides._fit_bullets(["очень длинная мысль " * 12], 3.0, 1.0))

# A long heading used to run three lines out of a 1.15in band, through the
# accent rule and across the first bullet. The geometry checks were blind to it:
# the BOX was the right size and only the words were too big for it.
_long = ("Очень длинный заголовок слайда, который раньше заезжал "
         "на фотографию справа")
check("a long heading is set smaller than a short one",
      slides._fit_heading(_long, 6.8, 1.15) < slides._fit_heading("Итоги", 6.8, 1.15),
      (slides._fit_heading(_long, 6.8, 1.15), slides._fit_heading("Итоги", 6.8, 1.15)))
check("a short heading keeps the full size",
      slides._fit_heading("Целевые уровни в крови", 6.8, 1.15) == 27)
check("even an absurd heading gets a size, not an exception",
      slides._fit_heading("я " * 200, 6.8, 1.15) >= 14)

_hd = slides.normalize_deck({
    "title": "T", "theme": "academic",
    "slides": [{"heading": _long,
                "bullets": ["Первая мысль с числом: 600 МЕ в сутки",
                            "Вторая мысль, тоже с числом: 125 мкг предел"],
                "image": "x"}]})
_outh = os.path.join(os.path.dirname(_out), "longhead.pptx")
slides.build_pptx(_hd, _outh, images={0: _imgp})
_ha = slides.pptx_audit(_outh)
check("...so nothing on the slide overflows its box", _ha["overflow"] == [],
      _ha["overflow"])
check("...and the heading still does not touch the photo",
      _ha["overlaps"] == [], _ha["overlaps"])

# And the overflow detector must be able to SEE one.
_ov = _s.shapes.add_textbox(_I(1.0), _I(1.0), _I(1.2), _I(0.3))
_ov.text_frame.text = "a great deal of text crammed into a very small box indeed"
_prs.save(_outb)
check("text too big for its box IS reported",
      slides.pptx_audit(_outb)["overflow"], "overflow detector is blind")

_tb3 = _s.shapes.add_textbox(_I(12.0), _I(7.0), _I(3.0), _I(1.0))
_tb3.text_frame.text = "hanging off the edge"
_prs.save(_outb)
check("a shape past the edge IS reported",
      slides.pptx_audit(_outb)["offslide"], "off-slide detector is blind")

print()
# "4 slides" means four slides in the FILE: the cover and the closing sources
# slide both count (live, 2026-09-12: a 4-slide request came back as 5) — but a
# tiny deck never spends a slide on sources: "3 слайда" came back as cover +
# ONE content slide + sources and was refused as a stub (journey 19). Below
# two content slides the sources move into the speaker notes instead.
_six = [{"heading": f"S{i}", "bullets": [f"fact {i}"]} for i in range(6)]
_src = [{"title": "a", "url": "https://example.com/a"}]
_d = slides.normalize_deck({"title": "T", "slides": list(_six), "sources": list(_src)}, "T", n_slides=4)
check("4 slides with sources = cover + 2 + sources", len(_d["slides"]) == 2 and not _d.get("sources_in_notes"), len(_d["slides"]))
_d = slides.normalize_deck({"title": "T", "slides": list(_six)}, "T", n_slides=4)
check("4 slides without sources = cover + 3", len(_d["slides"]) == 3, len(_d["slides"]))
_d = slides.normalize_deck({"title": "T", "slides": list(_six), "sources": list(_src)}, "T", n_slides=3)
check("3 slides with sources = cover + 2, sources into notes", len(_d["slides"]) == 2 and _d.get("sources_in_notes") and _d["sources"], _d)
_d = slides.normalize_deck({"title": "T", "slides": list(_six), "sources": list(_src)}, "T", n_slides=2)
check("never fewer than two content slides", len(_d["slides"]) == 2, len(_d["slides"]))
_d = slides.normalize_deck({"title": "T", "slides": list(_six), "sources": list(_src)}, "T", n_slides=6)
check("6 slides with sources = cover + 4 + sources", len(_d["slides"]) == 4 and not _d.get("sources_in_notes"), len(_d["slides"]))
import tempfile as _tf
_out = os.path.join(_tf.mkdtemp(), "tiny.pptx")
slides.build_pptx(slides.normalize_deck({"title": "T", "slides": list(_six), "sources": list(_src)}, "T", n_slides=3), _out)
from pptx import Presentation as _P
_prs = _P(_out)
check("the tiny deck file has exactly 3 slides", len(_prs.slides) == 3, len(_prs.slides))
check("its last slide carries the sources in the notes",
      "https://example.com/a" in _prs.slides[2].notes_slide.notes_text_frame.text)

print("%d/%d checks passed" % (_pass, _pass + _fail))
if _fail:
    sys.exit(1)

