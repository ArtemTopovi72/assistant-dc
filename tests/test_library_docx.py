"""DOCX: the help promised it, python-docx was not installed and the library refused
the extension, so a .docx got neither an inline answer nor an index entry.

Run: venv/Scripts/python.exe tests/test_library_docx.py
"""
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import docx
import library

d = docx.Document()
d.add_paragraph("Первый абзац про котов.")
t = d.add_table(rows=1, cols=2)
t.cell(0, 0).text, t.cell(0, 1).text = "Цена", "100 руб"
p = os.path.join(tempfile.mkdtemp(), "a.docx")
d.save(p)
txt = library.extract_text(p)
bad = 0
for name, ok in [(".docx is a supported library type", ".docx" in library.SUPPORTED_EXTS),
                 ("paragraph text extracted", "котов" in txt),
                 ("table cells extracted", "Цена | 100 руб" in txt)]:
    print(("PASS  " if ok else "FAIL  ") + name); bad += not ok
sys.exit(1 if bad else 0)
